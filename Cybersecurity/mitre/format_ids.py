import pandas as pd
from openpyxl import load_workbook

file_path = "c:/Users/crist/OneDrive/Desktop/mitre/mitre_attck_mapping_flow.xlsx"

wb = load_workbook(file_path)
ws = wb.active

data = [
    ["TA0001", "T1190", "G0087", "M1050, M1016", "DS0015", "DS0029"],
    ["TA0007", "T1082", "G0087", "M1026", "DS0015", "DS0029"],
    ["TA0009, TA0006", "T1005, T1003", "G0087", "M1057", "DS0015", "DS0029"]
]

for row_idx, row_data in enumerate(data, start=8):
    for col_idx, cell_value in enumerate(row_data, start=1):
        ws.cell(row=row_idx, column=col_idx, value=cell_value)

try:
    wb.save(file_path)
    print("Success: Updated mitre_attck_mapping_flow.xlsx")
except Exception as e:
    alt_path = "c:/Users/crist/OneDrive/Desktop/mitre/mitre_attck_mapping_flow_final.xlsx"
    wb.save(alt_path)
    print("Locked. Saved to: mitre_attck_mapping_flow_final.xlsx")
