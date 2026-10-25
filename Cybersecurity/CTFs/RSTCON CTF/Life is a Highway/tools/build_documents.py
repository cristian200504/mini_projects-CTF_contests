"""Convert the Markdown technical report to an editable, styled Word document.

This is a report-production utility, not a component of the CTF calculation.
Run with python-docx installed, then render the DOCX using a Word/PDF renderer.
"""
from pathlib import Path
import re
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parents[1]


def font(run, name, size=None, bold=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    run.bold = bold


def inline(paragraph, text, table=False):
    for token in re.split(r"(\*\*.*?\*\*|`[^`]+`)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`"):
            value = token[1:-1]
            if table and re.fullmatch(r"[0-9a-f]{64}", value):
                value = value[:32]+"\n"+value[32:]
            run = paragraph.add_run(value)
            font(run, "DejaVu Sans Mono", 8.7 if table else 9.4)
        else:
            paragraph.add_run(token)


def table_borders(table):
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = OxmlElement("w:"+name)
        for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
            element.set(qn("w:"+key), value)
        borders.append(element)
    props.append(borders)


def add_table(document, rows):
    rows = [[cell.strip() for cell in row] for row in rows]
    count = len(rows[0])
    if count == 2:
        widths = [2.15, 4.75]
    elif count == 4:
        widths = [1.75, 1.65, 1.65, 1.85]
    elif count == 5:
        widths = [0.7, 0.8, 2.2, 2.0, 1.2]
    else:
        widths = [6.9/count]*count
    if rows[0][0] == "Hex ID":
        widths = [0.8, 0.8, 2.25, 3.05]
    if rows[0][0] == "Parameter":
        widths = [1.65, 1.75, 3.5]
    if count == 3 and rows[0][0] != "Parameter":
        widths = [2.3]*3
    table = document.add_table(rows=1, cols=count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for index, values in enumerate(rows):
        cells = table.rows[0].cells if index == 0 else table.add_row().cells
        for cell, width, value in zip(cells, widths, values):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                margin = OxmlElement("w:"+side)
                margin.set(qn("w:w"), "85")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            props.append(margins)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E8EEF2" if index == 0 else ("F8FAFB" if index % 2 == 0 else "FFFFFF"))
            props.append(shading)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.03
            inline(p, value.strip(), table=True)
            for run in p.runs:
                if run.font.size is None:
                    run.font.size = Pt(9.6)
                if index == 0:
                    run.bold = True
            if count >= 4 and index > 0 and len(value) < 18:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        row_props = table.rows[index]._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        row_props.append(no_split)
        if index == 0:
            repeat = OxmlElement("w:tblHeader")
            row_props.append(repeat)
    table_borders(table)
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 0.4
    p.add_run().font.size = Pt(3)


def main():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = document.styles
    for name, size, bold in [("Normal", 11, False), ("Title", 25, False),
                              ("Heading 1", 15, True), ("Heading 2", 12, True),
                              ("Caption", 9.2, False)]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Title"].paragraph_format.space_after = Pt(12)
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(7)
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)
    styles["Caption"].paragraph_format.space_after = Pt(9)
    code = styles.add_style("CTF Code", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name, code.font.size = "DejaVu Sans Mono", Pt(8.8)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.space_after = Pt(7)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.left_indent = Inches(0.1)
    code.paragraph_format.right_indent = Inches(0.03)
    document.core_properties.title = "Life is a Highway CTF writeup"
    document.core_properties.subject = "CAN decoding and route reconstruction"
    document.core_properties.author = "ChatGPT"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Life is a Highway CTF writeup  |  ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    lines = (ROOT/"WRITEUP.md").read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith("```"):
            code_lines = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                code_lines.append(lines[index])
                index += 1
            p = document.add_paragraph(style="CTF Code")
            p.paragraph_format.keep_together = len(code_lines) < 12
            run = p.add_run("\n".join(code_lines))
            font(run, "DejaVu Sans Mono", 8.8)
            index += 1
            continue
        if line.startswith("#"):
            level = len(line)-len(line.lstrip("#"))
            text = line[level:].strip()
            if level == 1:
                document.add_paragraph(text, style="Title")
            else:
                document.add_paragraph(text, style=f"Heading {min(level-1, 2)}")
            index += 1
            continue
        if line.startswith("| "):
            rows = []
            while index < len(lines) and lines[index].startswith("|"):
                cells = lines[index].strip().strip("|").split("|")
                if not all(re.fullmatch(r"\s*:?-+:?\s*", cell) for cell in cells):
                    rows.append(cells)
                index += 1
            add_table(document, rows)
            continue
        match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if match:
            label, filename = match.groups()
            width = 6.2 if "route_map" in filename else (6.6 if "calibration" not in filename else 5.7)
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            image = p.add_run().add_picture(str(ROOT/filename), width=Inches(width))
            image._inline.docPr.set("descr", label)
            index += 1
            continue
        text_lines = [line]
        index += 1
        while index < len(lines) and lines[index].strip() and not lines[index].startswith(("#", "```", "|", "![")):
            text_lines.append(lines[index])
            index += 1
        text = " ".join(text_lines)
        p = document.add_paragraph(style="Caption" if text.startswith("Figure ") else "Normal")
        inline(p, text)
        if text.endswith(":"):
            p.paragraph_format.keep_with_next = True
    # The default Word template can carry a blue paragraph border in Title.
    # Remove inherited and direct paragraph borders, retaining table borders.
    for border in list(document.styles.element.xpath(".//w:pBdr")):
        border.getparent().remove(border)
    for border in list(document._element.xpath(".//w:pBdr")):
        border.getparent().remove(border)
    output = ROOT/"Life_is_a_Highway_Writeup.docx"
    document.save(output)
    print(output)


if __name__ == "__main__":
    main()
